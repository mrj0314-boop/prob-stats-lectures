"""2주차 원순열 슬라이드 생성. 양식은 _공용/deck.py, 그림은 ../애니메이션/."""
import math
import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

WEEK = Path(__file__).resolve().parents[1]
sys.dont_write_bytecode = True  # 폴더에 __pycache__ 만들지 않기
sys.path.insert(0, str(WEEK.parent / "_공용"))
from deck import GRAY, LIGHT, NAVY, ORANGE, WHITE, Deck, box, bullets, text  # noqa: E402

WOOD = (0xE3, 0xC5, 0x9E)


def ring(slide, cx, cy, r, labels, node=0.62, fill=NAVY, center_fill=None, link=None, highlight=()):
    """원형으로 배치한 노드(원). link 색이 있으면 이웃 노드를 선으로 연결."""
    from pptx.dml.color import RGBColor
    if center_fill:
        box(slide, Inches(cx - r * 0.72), Inches(cy - r * 0.72), Inches(r * 1.44), Inches(r * 1.44),
            RGBColor(*center_fill), MSO_SHAPE.OVAL)
    pts = [(cx + r * math.cos(math.pi / 2 - k * 2 * math.pi / len(labels)),
            cy - r * math.sin(math.pi / 2 - k * 2 * math.pi / len(labels))) for k in range(len(labels))]
    if link is not None:
        for k in range(len(pts)):
            (x1, y1), (x2, y2) = pts[k], pts[(k + 1) % len(pts)]
            ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
            ln.line.color.rgb = link
            ln.line.width = Pt(3)
    for (x, y), lab, k in zip(pts, labels, range(len(pts))):
        c = box(slide, Inches(x - node / 2), Inches(y - node / 2), Inches(node), Inches(node),
                ORANGE if k in highlight else fill, MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r_ = p.add_run()
        r_.text = lab
        from deck import style_run
        style_run(r_, 14 if len(lab) > 1 else 18, WHITE, True)


MEDIA = WEEK / "생성미디어"  # Google Flow로 만든 영상·이미지
d = Deck("확률과 통계 · 2주차 원순열  |  교과서 미래엔 Ⅰ-1 p.11~12", WEEK / "애니메이션")

s = d.title("확률과 통계", "2주차 · 원순열 (circular permutation)",
            "공과대학 1학년  |  교과서: 미래엔 「확률과 통계」 Ⅰ-1 순열과 조합 p.11~12",
            "고등학교에서 배운 순열 복습 후 원형 배열로 확장. 교과서 '생각 열기'(p.11)의 A, B, C, D 원탁 회의 상황으로 시작")
d.picture(s, MEDIA / "회전개념도.png", Inches(7.4), Inches(0.9), Inches(5.4), Inches(3.3))

d.video("도입: 모두 한 칸씩 옮겨 앉으면?", MEDIA / "도입영상.mp4", MEDIA / "도입영상_첫장면.png",
        "질문: 모두 옆자리로 옮겨 앉은 뒤의 배치는 처음과 다른 배치인가?",
        "8초 영상, 슬라이드가 열리면 자동 재생 (소리 있음). Google Flow(Veo 3.1 Lite)로 생성한 AI 영상: "
        "중간에 학생 수가 늘어나는 생성 오류가 있으므로 화면 속 인원수는 언급하지 않기. "
        "질문에 대한 답은 다음 슬라이드들에서 확인")

d.cards("오늘의 학습 목표", [
    ("01", "원순열의 뜻", "원탁 회의 비유"),
    ("02", "공식 (n − 1)!", "한 사람 고정하기"),
    ("03", "이웃 조건·다각형 탁자", "끈으로 묶기, 탁자 모양 비교"),
    ("04", "팔찌·목걸이 순열", "뒤집기 비유 (심화)"),
], "교과서 학습 목표(p.11): 원순열을 이해하고 그 수를 구할 수 있다. 04는 교과서 범위 밖 심화 내용")

d.content("원순열이란: 원탁 회의 비유", [
    "원순열(circular permutation)",
    ("서로 다른 대상을 원형으로 배열하는 순열", 1),
    "회전해서 겹치는 배열은 같은 것으로 간주",
    "비유: 원탁 회의",
    ("모두 한 칸씩 옮겨 앉아도 양옆 사람은 그대로", 1),
    ("누가 누구 옆에 앉았는지가 핵심", 1),
    "일렬 순열과의 차이: 맨 앞·맨 뒤 자리 없음",
], "RoundTableRotate.gif", "교과서 p.11 생각 열기와 같은 상황. 학생들에게 '자리를 다 바꿨는데 왜 같은 배치인가?' 질문. "
   "양옆 사람이 바뀌지 않았다는 답 유도")

d.gallery("생활 속 원순열", [
    (MEDIA / "원탁토론.png", "원탁 회의"),
    (MEDIA / "구절판.png", "정팔각형 구절판"),
    (MEDIA / "팔찌.png", "구슬 팔찌"),
], [
    "공통점: 맨 앞·맨 뒤 없이 둥글게 놓인 배열",
    "차이점: 팔찌는 뒤집을 수 있음 → 오늘 마지막 심화 주제",
], "그림마다 '돌려도 같은가?', '뒤집어도 같은가?'를 학생들에게 질문. 구절판은 교과서 p.12 문제 2와 연결")

d.content("일렬 순열과 원순열의 관계", [
    "일렬 배열 ABCD, BCDA, CDAB, DABC",
    ("맨 앞 사람을 맨 뒤로 보내는 회전 관계", 1),
    "원탁에 앉히면 모두 같은 배치",
    "원순열 1가지 ↔ 일렬 배열 n가지 대응",
    "원순열의 수 = n! ÷ n",
    ("4명: 4! ÷ 4 = 24 ÷ 4 = 6", 1),
], "LineToCircle.gif", "교과서 p.11의 4!/4 = 3! 유도와 같은 흐름. "
   "'중복해서 센 만큼 나눈다'는 아이디어 강조. 이후 팔찌 순열의 ÷2와 같은 원리")

d.content("공식: 한 사람을 고정하기", [
    "한 사람(A)의 자리를 기준으로 고정",
    ("회전으로 겹치는 경우가 저절로 사라짐", 1),
    "나머지 (n − 1)명의 일렬 배치와 같음",
    ("원순열의 수 = n! / n = (n − 1)!", "key"),
    "예: 4명 → 3! = 6",
    "예: 5명 → 4! = 24",
], "FixOnePerson.gif", "교과서는 n!/n = (n-1)! 한 가지로 유도(p.12). 한 사람 고정은 같은 결과를 다르게 보는 방법으로 보충. "
   "화면의 6개 원탁을 학생들이 직접 그려 보게 하기")

# 예제: 두 문제 카드 + 원탁 그림
s = d.frame("예제")
for i, (head, prob, sol, labels, hl) in enumerate([
    ("예제 1", "서로 다른 7명이 원탁에 둘러앉는 경우의 수",
     ["한 명 고정 → 나머지 6명 일렬 배치", "(7 − 1)! = 6! = 720"], ["A", "", "", "", "", "", ""], (0,)),
    ("예제 2", "6명이 원탁에 앉을 때 A와 B가 마주 보고 앉는 경우의 수",
     ["A 고정 → B는 맞은편 1자리로 결정", "나머지 4명: 4! = 24"], ["A", "", "", "B", "", ""], (0, 3)),
]):
    x = 0.6 + i * 6.25
    card = box(s, Inches(x), Inches(1.4), Inches(5.9), Inches(5.4), LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
    card.adjustments[0] = 0.04
    text(s, Inches(x + 0.3), Inches(1.55), Inches(2), Inches(0.5), head, 20, ORANGE, True)
    text(s, Inches(x + 0.3), Inches(2.05), Inches(5.3), Inches(0.9), prob, 19, NAVY, True)
    ring(s, x + 1.45, 4.75, 0.95, labels, node=0.48, fill=GRAY, center_fill=WOOD, highlight=hl)
    bullets(s, [sol[0], (sol[1], "key")], x=Inches(x + 2.75), y=Inches(3.55), w=Inches(3.0), h=Inches(2.8), size=17)
d.notes(s, "교과서 p.12 예제 1⑴, 문제 1⑵와 같은 유형 (인원 변경). "
           "예제 2: A를 고정하면 원순열의 회전 문제가 사라지고, B의 자리는 A 맞은편 하나로 결정. 일반 순열로 바뀌는 과정 강조")

d.content("이웃 조건: 끈으로 묶기", [
    "조건: 특정 두 사람 A, B가 반드시 이웃",
    "비유: 두 사람을 끈으로 묶어 한 사람처럼 취급",
    "① 묶음 1개 + 나머지 4명 = 5개의 원순열",
    ("(5 − 1)! = 24", 1),
    "② 묶음 안 순서: AB, BA",
    ("2! = 2", 1),
    ("24 × 2 = 48", "key"),
], "NeighborsGlue.gif", "교과서 p.12 예제 1⑵(회장·부회장 이웃)와 같은 방법. "
   "묶음 안 순서를 빠뜨리는 실수가 가장 흔함. 일렬 순열에서 배운 '이웃 조건'과 같은 방법임을 연결")

d.content("다각형 탁자: 탁자 모양 비교", [
    "정n각형 탁자, 각 변에 1명씩",
    ("한 칸씩 돌리면 모양 일치 → 원탁과 같음", 1),
    ("정팔각형 8명: (8 − 1)! = 7!", 1),
    "정사각형 탁자, 각 변에 2명씩 (8명)",
    ("한 칸 이동: 변의 왼쪽 끝 ↔ 오른쪽 끝, 다른 배치", 1),
    ("두 칸 이동(90° 회전): 같은 배치", 1),
    ("원순열 7! × 2 = 10080", "key"),
], "PolygonTable.gif", "교과서 p.12 문제 2(정팔각형 구절판), p.23 중단원 마무리 05(정팔각형 탁자)와 연결. "
   "정사각형 탁자는 교과서 밖 확장 유형이지만 '회전해서 겹치는가'라는 같은 기준으로 판단")

d.content("심화: 팔찌·목걸이 순열 (교과서 범위 밖)", [
    "팔찌, 목걸이: 뒤집어도 같은 물건",
    "시계 방향 배열 = 반시계 방향 배열",
    "원순열 2개 → 팔찌 1개 대응",
    ("팔찌 순열의 수 = (n − 1)! ÷ 2", "key"),
    ("n ≥ 3일 때 성립", 1),
    "예: 서로 다른 구슬 5개 → 24 ÷ 2 = 12",
], "NecklaceFlip.gif", "원탁은 뒤집을 수 없으므로 ÷2를 하지 않는다는 점 대비. n=3일 때 직접 그려서 확인: 2!/2 = 1")

# 공학 예시: 링 네트워크
s = d.frame("공학 속 원순열: 링 네트워크")
bullets(s, [
    "링(ring) 네트워크",
    ("컴퓨터를 원형으로 연결하는 통신 구조", 1),
    "노드 배치 순서의 경우의 수 = 원순열",
    "신호 방향을 구분하지 않으면 팔찌 순열",
    ("노드 6개: 5! ÷ 2 = 60가지", 1),
    "같은 원리: 원형 회전 테이블 공정, 터빈 블레이드 배치",
])
ring(s, 9.7, 4.1, 2.0, [f"PC{k}" for k in range(1, 7)], node=0.95, link=GRAY)
text(s, Inches(8.2), Inches(3.85), Inches(3), Inches(0.5), "신호 흐름", 16, GRAY, align=PP_ALIGN.CENTER)
d.notes(s, "토큰 링 네트워크 예시. 노드 사이 케이블 길이를 최소화하는 배치 문제 등에서 경우의 수 계산이 출발점")

d.flow("정리", [
    ("원순열", "(n − 1)!"), ("이웃 조건", "묶음 원순열 × 묶음 안 순서"),
    ("팔찌·목걸이", "(n − 1)! ÷ 2"), ("핵심 원리", "중복해서 센 만큼 나누기"),
], [
    "회전으로 같아지는 것 n개 → ÷ n,  뒤집어서 같아지는 것 2개 → ÷ 2",
    "막히면 한 사람 고정 후 일렬 순열로 바꿔 생각",
], "공식 세 개를 '중복 제거' 하나의 원리로 묶어 정리")

d.exercises("연습문제", [
    "1. 회장과 부회장을 포함한 7명이 원탁에 둘러앉을 때",
    ("⑴ 7명이 앉는 경우의 수   ⑵ 회장과 부회장이 이웃하게 앉는 경우의 수", 1),
    "2. 정육각형 모양의 탁자 각 변에 1명씩 6명이 둘러앉는 경우의 수",
    "3. 정삼각형 모양의 탁자 각 변에 2명씩 6명이 둘러앉는 경우의 수",
    "4. 남학생 3명, 여학생 3명이 원탁에 번갈아 앉는 경우의 수",
    ("남학생 먼저 원형 배치 후 사이 자리에 여학생 배치", 1),
    "5. (심화) 서로 다른 색 구슬 6개로 만들 수 있는 팔찌의 수",
], "정답: 1) ⑴ 6! = 720 ⑵ (6-1)! × 2! = 240  2) 5! = 120  3) 5! × 2 = 240  4) 2! × 3! = 12  5) 5!/2 = 60\n"
   "교과서 연계: 1번 ↔ p.12 예제 1, p.24 07 / 2번 ↔ p.23 05 / 3번 ↔ 다각형 탁자 확장")

d.save(WEEK / "02주차_원순열.pptx")
