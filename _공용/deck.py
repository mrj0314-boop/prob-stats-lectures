"""강의 슬라이드 공용 양식 (python-pptx).

양식: Beamer Metropolis 스타일 (남색 헤더 + 주황 포인트, 16:9)
그림: <주차>/애니메이션/*.gif (manim 애니메이션, 슬라이드쇼에서 자동 재생)
"""
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x23, 0x37, 0x3B)
ORANGE = RGBColor(0xEB, 0x81, 0x1B)
GRAY = RGBColor(0x8A, 0x94, 0x99)
LIGHT = RGBColor(0xE9, 0xED, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "맑은 고딕"


def no_word_split(p):
    """한글 단어가 줄 끝에서 둘로 잘리지 않도록 단어 단위 줄바꿈."""
    ppr = p._p.get_or_add_pPr()
    ppr.set("eaLnBrk", "0")
    ppr.set("latinLnBrk", "0")


def style_run(run, size, color=NAVY, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ko-KR")
    rpr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def box(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE, line=None):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(2.5)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    return shp


def text(slide, x, y, w, h, s, size, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    no_word_split(p)
    r = p.add_run()
    r.text = s
    style_run(r, size, color, bold)
    return tb


def bullets(slide, items, x=Inches(0.6), y=Inches(1.5), w=Inches(5.6), h=Inches(5.3), size=20):
    """items: 문자열 또는 (문자열, 수준) 튜플.
    수준 1 = 들여쓴 보조 설명, "key" = 주황 굵은 글씨로 강조한 핵심 공식."""
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        s, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6 if lvl == 1 else 14)
        no_word_split(p)
        # 줄바꿈된 둘째 줄이 글머리표 뒤 글자에 맞춰 시작하도록 내어쓰기
        mar, ind = {0: (0.32, -0.32), 1: (0.62, -0.26)}.get(lvl, (0, 0))
        p._p.get_or_add_pPr().set("marL", str(Inches(mar)))
        p._p.get_or_add_pPr().set("indent", str(Inches(ind)))
        if lvl == "key":
            r = p.add_run()
            r.text = s
            style_run(r, size + 4, ORANGE, True)
            continue
        mark = p.add_run()
        mark.text = "–  " if lvl else "■  "
        style_run(mark, size - 4 if lvl else size - 6, GRAY if lvl else ORANGE)
        r = p.add_run()
        r.text = s
        style_run(r, size - 3 if lvl else size, GRAY if lvl else NAVY)
    return tf


def arrow(slide, x1, y1, x2, y2, color=GRAY, width=3):
    a = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    a.line.color.rgb = color
    a.line.width = Pt(width)
    ln = a.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    return a


AUTOPLAY_TIMING = """<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>
<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/>
</p:cond></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>
<p:childTnLst><p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect">
<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr>
<p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst>
</p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/>
</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def autoplay(slide, spid):
    """동영상을 클릭 없이 슬라이드가 열릴 때 자동 재생하도록 타이밍 교체.
    (PowerPoint에서 '자동 실행'으로 설정했을 때 저장되는 구조와 같음)"""
    from lxml import etree
    old = slide._element.find(qn("p:timing"))
    if old is not None:
        slide._element.remove(old)
    slide._element.append(etree.fromstring(AUTOPLAY_TIMING.replace("{spid}", str(spid))))


class Deck:
    def __init__(self, footer, fig_dir):
        self.footer = footer
        self.fig_dir = fig_dir
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def _new(self):
        return self.prs.slides.add_slide(self.blank)

    @staticmethod
    def notes(slide, s):
        slide.notes_slide.notes_text_frame.text = s

    def title(self, title, subtitle, tag, note=""):
        s = self._new()
        box(s, 0, 0, self.prs.slide_width, self.prs.slide_height, NAVY)
        text(s, Inches(0.9), Inches(2.3), Inches(11), Inches(1.2), title, 48, WHITE, True)
        text(s, Inches(0.9), Inches(3.4), Inches(11), Inches(0.8), subtitle, 26, WHITE)
        box(s, Inches(0.9), Inches(4.35), Inches(5.5), Inches(0.05), ORANGE)
        text(s, Inches(0.9), Inches(4.6), Inches(11), Inches(0.6), tag, 16, LIGHT)
        self.notes(s, note)
        return s

    def frame(self, title):
        """헤더, 주황 줄, 바닥글이 있는 빈 본문 슬라이드."""
        s = self._new()
        w = self.prs.slide_width
        box(s, 0, 0, w, Inches(1.0), NAVY)
        box(s, 0, Inches(1.0), w, Inches(0.06), ORANGE)
        text(s, Inches(0.5), 0, Inches(12), Inches(1.0), title, 28, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(0.5), Inches(7.0), Inches(8), Inches(0.4), self.footer, 11, GRAY)
        text(s, Inches(11.8), Inches(7.0), Inches(1.0), Inches(0.4), str(len(self.prs.slides)), 11, GRAY,
             align=PP_ALIGN.RIGHT)
        return s

    @staticmethod
    def picture(slide, path, x, y, max_w, max_h):
        w, h = Image.open(path).size
        scale = min(max_w / w, max_h / h)
        pw, ph = int(w * scale), int(h * scale)
        return slide.shapes.add_picture(str(path), x + (max_w - pw) // 2, y + (max_h - ph) // 2, pw, ph)

    def content(self, title, items, gif, note=""):
        """왼쪽 글머리표 + 오른쪽 애니메이션."""
        s = self.frame(title)
        bullets(s, items)
        self.picture(s, self.fig_dir / gif, Inches(6.4), Inches(1.3), Inches(6.6), Inches(5.6))
        self.notes(s, note)
        return s

    def cards(self, title, cards, note=""):
        """번호 카드 4장: (번호, 제목, 설명)."""
        s = self.frame(title)
        for i, (num, head, sub) in enumerate(cards):
            x = Inches(0.6 + i * 3.1)
            card = box(s, x, Inches(1.9), Inches(2.85), Inches(3.9), LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE)
            card.adjustments[0] = 0.06
            text(s, x + Inches(0.3), Inches(2.2), Inches(2.3), Inches(0.9), num, 40, ORANGE, True)
            text(s, x + Inches(0.3), Inches(3.3), Inches(2.4), Inches(1.2), head, 22, NAVY, True)
            text(s, x + Inches(0.3), Inches(4.6), Inches(2.4), Inches(1.0), sub, 15, GRAY)
        self.notes(s, note)
        return s

    def flow(self, title, steps, tail, note=""):
        """화살표로 이은 개념 상자 4개 + 아래 글머리표."""
        s = self.frame(title)
        for i, (head, sub) in enumerate(steps):
            x = Inches(0.6 + i * 3.15)
            card = box(s, x, Inches(1.9), Inches(2.6), Inches(2.1), ORANGE if i == len(steps) - 1 else NAVY,
                       MSO_SHAPE.ROUNDED_RECTANGLE)
            card.adjustments[0] = 0.1
            text(s, x, Inches(2.15), Inches(2.6), Inches(0.7), head, 22, WHITE, True, PP_ALIGN.CENTER)
            text(s, x + Inches(0.15), Inches(2.9), Inches(2.3), Inches(1.0), sub, 16, LIGHT, align=PP_ALIGN.CENTER)
            if i < len(steps) - 1:
                arrow(s, x + Inches(2.65), Inches(2.95), x + Inches(3.1), Inches(2.95))
        bullets(s, tail, y=Inches(4.5), w=Inches(12), h=Inches(2.0))
        self.notes(s, note)
        return s

    def video(self, title, mp4, poster, question, note=""):
        """슬라이드 가운데 동영상 (슬라이드가 열리면 자동 재생) + 아래 질문 한 줄."""
        s = self.frame(title)
        w, h = Inches(8.8), Inches(4.95)
        x = (self.prs.slide_width - w) // 2
        movie = s.shapes.add_movie(str(mp4), x, Inches(1.3), w, h, poster_frame_image=str(poster),
                                   mime_type="video/mp4")
        autoplay(s, movie.shape_id)
        text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.5), question, 20, ORANGE, True, PP_ALIGN.CENTER)
        self.notes(s, note)
        return s

    def gallery(self, title, items, tail, note=""):
        """그림 3장을 나란히 + 그림마다 설명 한 줄 + 아래 글머리표."""
        s = self.frame(title)
        for i, (path, cap) in enumerate(items):
            x = Inches(0.6 + i * 4.1)
            self.picture(s, path, x, Inches(1.35), Inches(3.9), Inches(2.6))
            text(s, x, Inches(4.0), Inches(3.9), Inches(0.5), cap, 18, NAVY, True, PP_ALIGN.CENTER)
        bullets(s, tail, y=Inches(4.8), w=Inches(12), h=Inches(2.0))
        self.notes(s, note)
        return s

    def exercises(self, title, items, note=""):
        s = self.frame(title)
        bullets(s, items, w=Inches(12.2), size=21)
        self.notes(s, note)
        return s

    def save(self, path):
        self.prs.save(path)
        print("->", path)
